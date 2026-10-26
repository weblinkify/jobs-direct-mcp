"""Application documents built from the candidate's own text: tailored resume, cover letter,
and Word export. Nothing is invented: bullets come from the CV, re-ordered by relevance to the job."""

from __future__ import annotations

import re
from datetime import date
from pathlib import Path
from typing import Optional

from .ats import SKILL_LEXICON, has_term, candidate_years, extract_domains

_BULLET_RE = re.compile(r"^\s*(?:[-*•▪]|\d+[.)])\s+(.*\S)\s*$")


def cv_bullets(cv: str) -> list[str]:
    """Achievement lines from a CV: explicit bullets, or failing that, sentences with a number in them."""
    bullets = [m.group(1) for line in cv.splitlines() if (m := _BULLET_RE.match(line))]
    if bullets:
        return bullets
    sentences = re.split(r"(?<=[.!?])\s+", cv)
    return [s.strip() for s in sentences if re.search(r"\d", s) and len(s.split()) >= 5]


def job_keywords(job_description: str) -> list[str]:
    lower = job_description.lower()
    return [s for s in SKILL_LEXICON if has_term(s, lower)]


def rank_bullets(bullets: list[str], keywords: list[str]) -> list[tuple[str, list[str]]]:
    """Sorts bullets by how many job keywords they mention; ties keep CV order. Quantified bullets win ties."""
    scored = []
    for index, bullet in enumerate(bullets):
        lower = bullet.lower()
        hits = [k for k in keywords if has_term(k, lower)]
        has_metric = 1 if re.search(r"\d", bullet) else 0
        scored.append((-len(hits), -has_metric, index, bullet, hits))
    scored.sort()
    return [(bullet, hits) for _, _, _, bullet, hits in scored]


def tailor_resume(cv: str, job_description: str, role: Optional[str] = None,
                  company: Optional[str] = None) -> str:
    role = role or "the role"
    company = company or "the company"
    keywords = job_keywords(job_description)
    cv_lower = cv.lower()
    have = [k for k in keywords if has_term(k, cv_lower)]
    gaps = [k for k in keywords if k not in have]
    domains = list(extract_domains(cv))
    years = candidate_years(cv)
    ranked = rank_bullets(cv_bullets(cv), keywords)

    years_txt = f"{years}+ years of experience" if years else "Hands-on experience"
    domain_txt = f" in {', '.join(domains[:2])}" if domains else ""
    skills_txt = f", strongest in {', '.join(have[:5])}" if have else ""

    lines = [f"# Tailored resume: {role} at {company}", "", "## Professional summary",
             f"{years_txt}{domain_txt}{skills_txt}. Targeting {role} at {company}.", "",
             "## Keyword alignment",
             f"- **Matched**: {', '.join(have) if have else 'none of the tracked keywords'}",
             f"- **Missing**: {', '.join(gaps) if gaps else 'none'}", "",
             "## Experience bullets, most relevant first"]
    if ranked:
        for bullet, hits in ranked[:8]:
            tag = f"  _(matches: {', '.join(hits)})_" if hits else ""
            lines.append(f"- {bullet}{tag}")
    else:
        lines.append("- No bullet points found in the CV. Add achievements as '- ' lines to rank them.")
    weak = [b for b, _ in ranked if not re.search(r"\d", b)]
    if weak:
        lines += ["", "## Add a number to these (STAR: Situation, Task, Action, Result)"]
        lines += [f"- {b}" for b in weak[:4]]
    if gaps:
        lines += ["", "## Gaps to address honestly"]
        lines += [f"- {g}: mention adjacent experience or a plan to ramp up." for g in gaps[:5]]
    return "\n".join(lines)


def cover_letter(cv: str, job_description: str, company: str, role: str,
                 candidate_name: Optional[str] = None, hiring_manager: Optional[str] = None,
                 today: Optional[date] = None) -> str:
    today = today or date.today()
    keywords = job_keywords(job_description)
    ranked = rank_bullets(cv_bullets(cv), keywords)
    proof = [b.rstrip(".") for b, _ in ranked[:2]]
    have = [k for k in keywords if has_term(k, cv.lower())]
    greeting = f"Dear {hiring_manager}," if hiring_manager else f"Dear {company} hiring team,"
    skills = f" My background in {', '.join(have[:3])} lines up with what you are asking for." if have else ""

    paragraphs = [
        today.strftime("%d %B %Y"),
        greeting,
        f"I am applying for the {role} position at {company}.{skills}",
    ]
    if proof:
        evidence = "Two examples of the work I would bring:\n" + "\n".join(f"- {p}." for p in proof)
        paragraphs.append(evidence)
    paragraphs += [
        f"I would welcome a conversation about how I can contribute to {company}. Thank you for your time.",
        f"Kind regards,\n{candidate_name or '[Your name]'}",
    ]
    return "\n\n".join(paragraphs)


def safe_filename(name: str, suffix: str) -> str:
    stem = re.sub(r"[^A-Za-z0-9._-]+", "_", name.strip()).strip("._") or "document"
    return stem[:80] + suffix


def export_docx(title: str, markdown: str, out_dir: Path, filename: Optional[str] = None) -> Path:
    """Writes simple Markdown (headings, bullets, **bold**: lead-ins, code fences) to a .docx file."""
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    heading = doc.add_heading(title, level=0)
    heading.runs[0].font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    in_code = False
    for raw in markdown.splitlines():
        line = raw.rstrip()
        if line.strip().startswith("```"):
            in_code = not in_code
            continue
        if in_code:
            run = doc.add_paragraph().add_run(line)
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            continue
        stripped = line.strip()
        if not stripped:
            continue
        level = len(stripped) - len(stripped.lstrip("#"))
        if 1 <= level <= 3 and stripped[level:level + 1] == " ":
            doc.add_heading(stripped[level + 1:], level=level)
        elif stripped[:2] in ("- ", "* "):
            _add_rich(doc.add_paragraph(style="List Bullet"), stripped[2:])
        else:
            _add_rich(doc.add_paragraph(), stripped)

    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / safe_filename(Path(filename).stem if filename else title, ".docx")
    doc.save(path)
    return path


def _add_rich(paragraph, text: str) -> None:
    """Adds text to a paragraph, rendering **bold** and _italic_ spans."""
    for part in re.split(r"(\*\*[^*]+\*\*|_[^_]+_)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("_") and part.endswith("_") and len(part) > 2:
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)
